"""Build a Spanish technical report PDF from the experiment workbook."""

from __future__ import annotations

import math
import statistics
import tempfile
from io import BytesIO
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
WORKBOOK = ROOT / "results" / "experimental_results.xlsx"
OUTPUT = ROOT / "reporte_tecnico.pdf"

ALGORITHMS = ("Backtracking", "SA", "GA")
PROBLEM_LABELS = {
    "N-Queens": "N-Reinas",
    "Knapsack 0/1": "Mochila 0/1",
    "Graph Coloring": "Coloreado de grafos",
    "Eggholder": "Eggholder",
}


def load_results(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["results"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    return [dict(zip(headers, row)) for row in rows]


def mean(values: list[float]) -> float:
    return statistics.fmean(values) if values else math.nan


def sample_std(values: list[float]) -> float:
    return statistics.stdev(values) if len(values) > 1 else 0.0


def grouped_results(rows: list[dict[str, Any]]) -> dict[tuple[str, str, str], list[dict[str, Any]]]:
    groups: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        groups[(row["problem"], str(row["instance"]), row["algorithm"])].append(row)
    return groups


def format_objective(group: list[dict[str, Any]]) -> str:
    values = [float(row["objective"]) for row in group if row["objective"] is not None]
    if not values:
        return "Sin resultado"
    return f"{mean(values):.2f} ± {sample_std(values):.2f}"


def chart_results(rows: list[dict[str, Any]], path: Path) -> None:
    groups = grouped_results(rows)
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.size": 9,
        "axes.edgecolor": "#555555",
        "axes.linewidth": 0.7,
        "xtick.color": "#333333",
        "ytick.color": "#333333",
    })
    fig, axes = plt.subplots(2, 2, figsize=(9, 6.5))
    panels = [
        ("N-Queens", ["8", "50", "100"], "Conflictos promedio"),
        ("Knapsack 0/1", ["30"], "Beneficio promedio"),
        ("Graph Coloring", ["50", "1000"], "Conflictos promedio"),
        ("Eggholder", ["2D"], "Valor de la función"),
    ]
    line_styles = {
        "Backtracking": ("#222222", "o", "-"),
        "SA": ("#777777", "s", "--"),
        "GA": ("#aaaaaa", "^", ":"),
    }
    for ax, (problem, instances, ylabel) in zip(axes.flat, panels):
        x = list(range(len(instances)))
        algorithms = ("Backtracking", "SA", "GA") if problem != "Eggholder" else ("SA", "GA")
        for algorithm in algorithms:
            means: list[float] = []
            deviations: list[float] = []
            for instance in instances:
                group = groups.get((problem, instance, algorithm), [])
                values = [float(row["objective"]) for row in group if row["objective"] is not None]
                means.append(mean(values) if values else math.nan)
                deviations.append(sample_std(values) if values else 0.0)
            color, marker, line_style = line_styles[algorithm]
            ax.errorbar(
                x, means, yerr=deviations, color=color, marker=marker,
                linestyle=line_style, linewidth=1.2, markersize=4,
                capsize=2, label=algorithm,
            )
        ax.set_xticks(x, instances)
        ax.set_xlabel("Tamaño")
        ax.set_ylabel(ylabel)
        ax.set_title(PROBLEM_LABELS[problem])
        ax.grid(axis="y", color="#dddddd", linewidth=0.5)
        ax.spines[["top", "right"]].set_visible(False)
        ax.legend(frameon=False, fontsize=8)
        if problem == "Eggholder":
            ax.axhline(-959.6407, color="#555555", linestyle="--", linewidth=0.8)
    fig.tight_layout()
    fig.savefig(path, dpi=150, bbox_inches="tight")
    plt.close(fig)


def make_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle", parent=styles["Title"], fontName="Helvetica-Bold",
        fontSize=20, leading=24, textColor=colors.black, spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="Section", parent=styles["Heading1"], fontName="Helvetica-Bold",
        fontSize=13, leading=16, textColor=colors.black, spaceBefore=9, spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="Subsection", parent=styles["Heading2"], fontName="Helvetica-Bold",
        fontSize=10, leading=13, textColor=colors.black, spaceBefore=6, spaceAfter=3,
    ))
    styles.add(ParagraphStyle(
        name="BodyCompact", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=9, leading=12, spaceAfter=5,
    ))
    styles.add(ParagraphStyle(
        name="Small", parent=styles["BodyText"], fontName="Helvetica",
        fontSize=7.5, leading=10, spaceAfter=4,
    ))
    return styles


def make_table(data: list[list[Any]], widths: list[float] | None = None) -> Table:
    table = Table(data, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("TEXTCOLOR", (0, 0), (-1, -1), colors.black),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("LEADING", (0, 0), (-1, -1), 9),
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
        ("LINEBELOW", (0, 1), (-1, -1), 0.25, colors.HexColor("#bbbbbb")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    return table


def page_footer(canvas, document) -> None:
    canvas.saveState()
    width, _ = letter
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.black)
    canvas.drawRightString(width - 0.68 * inch, 0.4 * inch, str(document.page))
    canvas.restoreState()


def build_report() -> Path:
    if not WORKBOOK.exists():
        raise FileNotFoundError(f"No se encontró el Excel experimental: {WORKBOOK}")
    rows = load_results(WORKBOOK)
    if not rows:
        raise ValueError("El Excel no contiene ejecuciones en la hoja results.")
    groups = grouped_results(rows)
    styles = make_styles()
    story: list[Any] = []

    story.extend([
        Paragraph("Búsqueda y optimización", styles["ReportTitle"]),
        Paragraph(
            f"Práctica · {date.today().strftime('%d/%m/%Y')}<br/>"
            "Nombre: ____________________ &nbsp;&nbsp; Grupo: ____________________",
            styles["BodyCompact"],
        ),
        Paragraph(
            "En este trabajo comparé backtracking, recocido simulado y algoritmo genético con "
            "varios problemas. Abajo resumo qué funcionó en estas pruebas y qué quedó pendiente.",
            styles["BodyCompact"],
        ),
        Paragraph("1. Resumen ejecutivo", styles["Section"]),
        Paragraph(
            f"Realicé {len(rows)} ejecuciones. Para los algoritmos aleatorios usé cinco semillas por "
            "caso y guardé los resultados en el Excel, junto con los tiempos, evaluaciones y curvas "
            "de convergencia. No hay un método que gane en todo: el resultado cambia según el problema.",
            styles["BodyCompact"],
        ),
        Paragraph(
            "<b>Nota sobre los datos:</b> no tenía la instancia de Mochila de la tarea anterior ni "
            "los archivos oficiales de grafos. Para poder probar el programa utilicé datos sintéticos. "
            "Los resultados de esos dos problemas son solo una prueba del funcionamiento y habrá que "
            "actualizarlos con las instancias de la materia antes de entregar.",
            styles["BodyCompact"],
        ),
        Paragraph("2. Problemas y representaciones", styles["Section"]),
        make_table([
            ["Problema", "Representación y objetivo"],
            ["N-Reinas", "Permutación de columnas por fila; minimizar pares de reinas en conflicto."],
            ["Mochila 0/1", "Vector binario de selección; maximizar beneficio sujeto a capacidad."],
            ["Coloreado", "Vector de color por nodo, cuatro colores; minimizar aristas con extremos del mismo color."],
            ["Eggholder", "Punto real (x, y) en [-512, 512]^2; minimizar la función Eggholder."],
        ], [1.35 * inch, 5.05 * inch]),
        Paragraph("3. Algoritmos", styles["Section"]),
        Paragraph(
            "<b>Backtracking.</b> N-Reinas aplica propagación mediante columnas/diagonales disponibles, "
            "elección de fila con dominio mínimo y orden de valores por menor impacto diagonal. "
            "Mochila ordena por beneficio/peso y usa una cota superior fraccional para podar. "
            "Coloreado ordena estáticamente los nodos por grado y descarta colores ya asignados a vecinos. "
            "En coloreado puse un límite de 200 000 nodos explorados.",
            styles["BodyCompact"],
        ),
        Paragraph(
            "<b>Recocido simulado.</b> Acepta mejoras y puede aceptar empeoramientos con probabilidad "
            "exp(-delta/T); el enfriamiento geométrico permite explorar al inicio y refinar después. "
            "El vecindario intercambia reinas, cambia el color de un nodo, modifica y repara una mochila, "
            "o perturba Eggholder con ruido gaussiano.",
            styles["BodyCompact"],
        ),
        Paragraph(
            "<b>Algoritmo genético.</b> Usa selección por torneo entre los cinco mejores, elitismo "
            "de dos individuos, cruza y mutación con tasa 0.2. La representación y reparación se "
            "adaptan a cada problema: permutación para reinas, bits con reparación de capacidad para "
            "mochila, colores por nodo y coordenadas reales para Eggholder.",
            styles["BodyCompact"],
        ),
        PageBreak(),
    ])

    story.extend([
        Paragraph("4. Configuración experimental", styles["Section"]),
        make_table([
            ["Elemento", "Configuración utilizada"],
            ["Ejecuciones", "5 por combinación; semillas reproducibles guardadas en Excel."],
            ["N-Reinas", "N = 8, 50, 100; SA: 2 000 iteraciones; GA: 200 generaciones, población 50."],
            ["Mochila", "30 objetos sintéticos (semilla de instancia 2026); capacidad = 35% del peso total."],
            ["Coloreado", "Grafos aleatorios de 50/1000 nodos; grado medio objetivo 4; 4 colores."],
            ["Eggholder", "Dominio [-512, 512]^2; SA: 2 000 iteraciones; GA: 200 generaciones, población 50."],
            ["Cómputo", "SA: 2 001 evaluaciones. GA: 9 650 evaluaciones contadas por ejecución."],
        ], [1.35 * inch, 5.05 * inch]),
        Spacer(1, 8),
        Paragraph(
            "El contador de backtracking mide nodos explorados; el de SA y GA, evaluaciones de la función "
            "objetivo. No es exactamente la misma unidad de trabajo, así que también comparé los tiempos. "
            "En GA se cuentan los individuos iniciales y los descendientes evaluados.",
            styles["BodyCompact"],
        ),
        Paragraph("5. Resultados cuantitativos", styles["Section"]),
        Paragraph("5.1 N-Reinas: conflictos promedio ± desviación estándar; factibilidad en 5 corridas.",
                  styles["Subsection"]),
    ])

    queen_table = [["N", "Algoritmo", "Conflictos", "Factibles", "Tiempo medio (s)"]]
    for n in ("8", "50", "100"):
        for algorithm in ALGORITHMS:
            group = groups[( "N-Queens", n, algorithm)]
            feasible = sum(bool(row["feasible"]) for row in group)
            queen_table.append([
                n, algorithm, format_objective(group), f"{feasible}/5",
                f"{mean([float(row['seconds']) for row in group]):.3f}",
            ])
    story.extend([
        make_table(queen_table, [0.6 * inch, 1.25 * inch, 1.35 * inch, 0.85 * inch, 1.5 * inch]),
        Spacer(1, 5),
        Paragraph(
            "Backtracking encontró solución en las cinco corridas para los tres tamaños. SA resolvió "
            "N=8 en las cinco corridas, pero no llegó a cero conflictos para N=50/100 con 2 000 pasos. "
            "GA redujo los conflictos en tamaños mayores, pero no encontró solución factible en esas "
            "corridas. En N=100 el tiempo de backtracking varió bastante entre semillas: la corrida más "
            "lenta tardó 17.27 s y exploró 400 438 nodos, aunque encontró solución.",
            styles["BodyCompact"],
        ),
        Paragraph("5.2 Mochila 0/1: beneficio (mayor es mejor).", styles["Subsection"]),
    ])
    knapsack_table = [["Algoritmo", "Beneficio medio ± DE", "Factibles", "Tiempo medio (s)"]]
    for algorithm in ALGORITHMS:
        group = groups[("Knapsack 0/1", "30", algorithm)]
        knapsack_table.append([
            algorithm, format_objective(group),
            f"{sum(bool(row['feasible']) for row in group)}/5",
            f"{mean([float(row['seconds']) for row in group]):.4f}",
        ])
    story.extend([
        make_table(knapsack_table, [1.25 * inch, 1.8 * inch, 0.9 * inch, 1.55 * inch]),
        Paragraph(
            "Backtracking obtuvo beneficio 558 en las cinco corridas y evaluó 152 nodos. GA también "
            "alcanzó 558 en las cinco corridas; SA obtuvo entre 555 y 558. El valor 558 es óptimo para "
            "esta instancia sintética según la búsqueda completa con cota. Esto no permite anticipar "
            "qué pasará con los datos oficiales.",
            styles["BodyCompact"],
        ),
        PageBreak(),
    ])

    story.extend([
        Paragraph("5.3 Coloreado de grafos: conflictos (menos es mejor).", styles["Subsection"]),
    ])
    coloring_table = [["Nodos", "Algoritmo", "Conflictos medios ± DE*", "Factibles", "Tiempo medio (s)"]]
    for n in ("50", "1000"):
        for algorithm in ALGORITHMS:
            group = groups[("Graph Coloring", n, algorithm)]
            feasible = sum(bool(row["feasible"]) for row in group)
            coloring_table.append([
                n, algorithm, format_objective(group), f"{feasible}/5",
                f"{mean([float(row['seconds']) for row in group]):.3f}",
            ])
    story.extend([
        make_table(coloring_table, [0.65 * inch, 1.15 * inch, 1.65 * inch, 0.85 * inch, 1.45 * inch]),
        Paragraph(
            "* En backtracking con 1000 nodos, tres de cinco ejecuciones alcanzaron el límite de 200 000 "
            "nodos; el promedio de cero conflictos refleja solo las dos corridas completas y no es "
            "comparable con los promedios de SA/GA. Esas corridas quedaron marcadas como incompletas, "
            "no como demostración de inexistencia de coloreo.",
            styles["Small"],
        ),
        Paragraph(
            "En el grafo de 50 nodos, backtracking coloreó sin conflictos en todas las corridas. GA "
            "fue factible en dos de cinco; SA no lo fue en ninguna con los parámetros usados. En 1000 "
            "nodos, GA obtuvo menos conflictos promedio que SA (270.2 frente a 368.2), pero ninguno "
            "de los dos encontró una asignación sin conflictos dentro del presupuesto.",
            styles["BodyCompact"],
        ),
        Paragraph("5.4 Eggholder: minimización.", styles["Subsection"]),
    ])
    egg_table = [["Algoritmo", "Objetivo medio ± DE", "Mejor corrida", "Tiempo medio (s)"]]
    for algorithm in ("SA", "GA"):
        group = groups[("Eggholder", "2D", algorithm)]
        values = [float(row["objective"]) for row in group]
        egg_table.append([
            algorithm, format_objective(group), f"{min(values):.2f}",
            f"{mean([float(row['seconds']) for row in group]):.4f}",
        ])
    story.extend([
        make_table(egg_table, [1.2 * inch, 1.8 * inch, 1.25 * inch, 1.55 * inch]),
        Paragraph(
            "GA encontró el mejor valor de las corridas (-959.61), muy cercano al mínimo global "
            "conocido (-959.64), y superó a SA en promedio. Sin embargo, su desviación estándar fue "
            "mayor, mostrando variabilidad entre semillas; no garantiza hallar el óptimo global en "
            "cada ejecución.",
            styles["BodyCompact"],
        ),
        Spacer(1, 6),
    ])

    with tempfile.TemporaryDirectory() as temp_dir:
        chart_path = Path(temp_dir) / "summary.png"
        chart_results(rows, chart_path)
        story.extend([
            Image(BytesIO(chart_path.read_bytes()), width=6.45 * inch, height=4.52 * inch),
            Paragraph(
                "Cada punto muestra la media y las líneas verticales la desviación estándar. En coloreado "
                "de 1000 nodos, backtracking incluye solo corridas completadas; véase la nota "
                "de factibilidad y terminación.",
                styles["Small"],
            ),
            PageBreak(),
        ])

    story.extend([
        Paragraph("6. Discusión", styles["Section"]),
        Paragraph(
            "<b>Calidad y rapidez.</b> En estas pruebas, backtracking funcionó bien para N-Reinas y "
            "Mochila. En problemas más grandes el costo depende mucho de cuánto se pueda podar. SA hizo "
            "2 001 evaluaciones por corrida y GA hizo 9 650; GA consiguió mejores resultados en algunos "
            "casos, pero también tomó más tiempo.",
            styles["BodyCompact"],
        ),
        Paragraph(
            "<b>Representación.</b> La permutación de N-Reinas impide columnas repetidas desde el inicio; "
            "aun así, una permutación puede conservar conflictos diagonales. En cambio, coloreado y "
            "Mochila usan asignaciones que pueden violar restricciones, por lo que el objetivo penaliza "
            "conflictos o la reparación restaura la capacidad. La representación real de Eggholder permite "
            "explorar el dominio continuo, pero la superficie multimodal hace depender el resultado de "
            "semilla, perturbación y enfriamiento.",
            styles["BodyCompact"],
        ),
        Paragraph(
            "<b>Escalabilidad y límites.</b> Backtracking con cuatro colores puede completar rápidamente "
            "algunos grafos aleatorios dispersos de 1000 nodos, pero otros consumieron el presupuesto. "
            "El límite computacional se debe distinguir de una prueba de inviabilidad. Los resultados "
            "de cinco semillas son evidencia inicial, no una caracterización universal de cada algoritmo.",
            styles["BodyCompact"],
        ),
        Paragraph("7. Conclusiones", styles["Section"]),
        Paragraph(
            "Con estos datos no puedo decir que haya un ganador para todos los problemas. Backtracking "
            "fue efectivo cuando las restricciones y las cotas ayudaron a reducir la búsqueda. Los "
            "métodos aleatorios dieron alternativas para espacios grandes y Eggholder, aunque variaron "
            "dependiente de los parámetros. GA obtuvo el mejor valor observado para Eggholder y menos "
            "conflictos medios que SA en coloreado de 1000 nodos; SA no alcanzó factibilidad para N=50/100 "
            "en N-Reinas ni para los grafos evaluados. Estas conclusiones describen solo las instancias y "
            "presupuestos documentados.",
            styles["BodyCompact"],
        ),
        Paragraph("8. Limitaciones y pasos antes de la entrega final", styles["Section"]),
        Paragraph(
            "1. Reemplazar la Mochila sintética por los pesos, beneficios y capacidad de la tarea anterior. "
            "2. Cargar los archivos oficiales de grafos, confirmar su número de nodos/aristas y reutilizar "
            "exactamente cada instancia entre algoritmos. 3. Repetir las ejecuciones y regenerar el Excel, "
            "las gráficas y este PDF. 4. Reportar parámetros y semillas efectivamente usados. "
            "5. Compartir el notebook Colab y el repositorio GitHub junto con el PDF.",
            styles["BodyCompact"],
        ),
        Paragraph("9. Artefactos reproducibles", styles["Section"]),
        Paragraph(
            "Código y notebook: optimizer_lab.py y run_experiments.ipynb. Resultados: "
            "results/experimental_results.xlsx (hojas results, convergence y summary), además de las "
            "gráficas PNG. Pruebas automatizadas: test_optimizer_lab.py. Comando utilizado: "
            "python optimizer_lab.py --runs 5 --iterations 2000 --generations 200 --population 50.",
            styles["BodyCompact"],
        ),
    ])

    doc = SimpleDocTemplate(
        str(OUTPUT), pagesize=letter, rightMargin=0.68 * inch, leftMargin=0.68 * inch,
        topMargin=0.65 * inch, bottomMargin=0.72 * inch,
        title="Reporte técnico: búsqueda y optimización",
        subject="Análisis de backtracking, recocido simulado y algoritmo genético",
    )
    doc.build(story, onFirstPage=page_footer, onLaterPages=page_footer)
    return OUTPUT


if __name__ == "__main__":
    print(f"PDF generado: {build_report()}")
